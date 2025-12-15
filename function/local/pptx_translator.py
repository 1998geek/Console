# 文件名: pptx_translator.py

import io
import re
import operator
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed

from pptx import Presentation
from pptx.text.text import _Run, _Paragraph

# =================================================================
#  请确保这些自定义模块在您的项目路径中
# =================================================================
from function.AzureAIClient import AzureAiClient
from app_config.keys_config import AZURE_OPENAI_ENDPOINT, AZURE_OPENAI_TOKEN, XINFERENCE_TRANSLATE_MODEL_NAME
from function.local.XinferenceChatClient import XinferenceChatClient

# =================================================================
#  1. 翻译函数 (新增 Xinference 支持)
# =================================================================

def get_llm_translation_xinference(text_to_translate: str, context: str, target_language: str) -> str:
    """调用通过 Xinference 部署的大模型进行翻译。"""
    try:
        xin_client = XinferenceChatClient()
        translated_text = f"*******[Xinference] 翻译失败: {text_to_translate}*******"

        messages = [
            {"role": "system", "content": f"你是一个专业的翻译助手。请将用户提供的文本翻译成{target_language}。请注意保持翻译的专业性和准确性，并参考上下文, 不要丢失标点符号以及空格。"},
            {"role": "user", "content": f"为了帮助你理解，这是该文本所在的完整段落上下文：\n---\n{context}\n---"},
            {"role": "user", "content": f"请仅翻译以下内容，不要添加任何额外的解释或标签：\n---\n{text_to_translate}\n---"}
        ]

        non_stream_response = xin_client.chat(
            messages=messages,
            model=XINFERENCE_TRANSLATE_MODEL_NAME,
            stream=False,
        )
        if non_stream_response and non_stream_response.choices:
            translated_text = non_stream_response.choices[0].message.content.strip()
            # 移除一些模型可能添加的思考过程标签
            cleaned_response = re.sub(r"<think>.*?</think>\n\n", "", translated_text, flags=re.DOTALL).strip()
            return cleaned_response

    except Exception as e:
        print(f"!!! 调用 Xinference 时出错: {e}")
        return translated_text
    
    return translated_text

def get_llm_translation(text_to_translate: str, context: str, target_language: str) -> str:
    """调用大模型进行翻译 (Azure 版本)"""
    az_AiClient = AzureAiClient(
        azure_endpoint=AZURE_OPENAI_ENDPOINT,
        api_key=AZURE_OPENAI_TOKEN
    )
    print(f"--- [Azure] 正在翻译 (原文): {text_to_translate}")
    translated_text = f"*******[Azure] 翻译失败: {text_to_translate}*******"
    messages = [
        {"role": "system", "content": f"你是一个专业的翻译助手。请将用户提供的文本翻译成{target_language}。请注意保持翻译的专业性和准确性，并参考上下文, 不要丢失标点符号。"},
        {"role": "user", "content": f"为了帮助你理解，这是该文本所在的完整段落上下文：\n---\n{context}\n---"},
        {"role": "user", "content": f"请仅翻译以下内容，不要添加任何额外的解释或标签：\n---\n{text_to_translate}\n---"}
    ]
    try:
        non_stream_response = az_AiClient.get_chat_completion(
            messages=messages,
            model="gpt-4.1-nano",
            stream=False,
        )
        if non_stream_response and non_stream_response.choices:
            return non_stream_response.choices[0].message.content.strip()
    except Exception as e:
        print(f"!!! 调用Azure AI时出错: {e}")
        return translated_text
    return translated_text

# =================================================================
#  2. 样式处理函数 (与之前版本保持一致并优化)
# =================================================================

def get_pptx_run_style(run: _Run):
    """获取一个PPTX run的关键格式信息。"""
    font = run.font
    color_info = None
    if font.color.type is not None:
        if font.color.type == 1: # MSO_COLOR_TYPE.RGB
            color_info = ('RGB', font.color.rgb)
        elif font.color.type == 2: # MSO_COLOR_TYPE.SCHEME
            color_info = ('SCHEME', font.color.theme_color, font.color.brightness)

    return (
        font.name, font.size, font.bold, font.italic, font.underline, color_info
    )

def apply_style_to_pptx_run(style_run: _Run, target_run: _Run):
    """将源run的格式应用到目标run。"""
    s_font = style_run.font
    t_font = target_run.font
    
    t_font.name = s_font.name
    t_font.size = s_font.size
    t_font.bold = s_font.bold
    t_font.italic = s_font.italic
    t_font.underline = s_font.underline

    s_color = s_font.color
    t_color = t_font.color
    if s_color.type is not None:
        if s_color.type == 1: # MSO_COLOR_TYPE.RGB
            t_color.rgb = s_color.rgb
        elif s_color.type == 2: # MSO_COLOR_TYPE.SCHEME
            t_color.theme_color = s_color.theme_color
            t_color.brightness = s_color.brightness

# =================================================================
#  3. 主翻译函数 (已重构为多线程模式)
# =================================================================

def translate_pptx_in_memory(
    input_file_obj, 
    target_language: str = "English", 
    progress_callback=None
):
    """
    主函数：使用多线程并发翻译PPTX文件流。
    采用“收集-执行-重建”模式，以确保稳定性和格式保真度。
    """
    prs = Presentation(input_file_obj)
    
    # --- 阶段 1: 收集所有需要翻译的任务 ---
    tasks = []
    all_paragraphs = []

    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                all_paragraphs.extend(shape.text_frame.paragraphs)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        all_paragraphs.extend(cell.text_frame.paragraphs)

    print(f"发现 {len(all_paragraphs)} 个段落。正在分析并拆分文本块...")

    for para in all_paragraphs:
        paragraph_context = para.text
        if not para.runs or not paragraph_context.strip():
            continue
        
        chunk_index_in_para = 0
        current_text_chunk = ""
        # 确保段落有内容再访问第一个run
        if not para.runs: continue
        style_of_chunk = para.runs[0]

        for run in para.runs:
            if get_pptx_run_style(run) == get_pptx_run_style(style_of_chunk):
                current_text_chunk += run.text
            else:
                if current_text_chunk.strip():
                    tasks.append({
                        'text': current_text_chunk,
                        'style_run': style_of_chunk,
                        'parent_para': para,
                        'context': paragraph_context,
                        'chunk_index': chunk_index_in_para
                    })
                    chunk_index_in_para += 1
                current_text_chunk = run.text
                style_of_chunk = run
        
        if current_text_chunk.strip():
            tasks.append({
                'text': current_text_chunk,
                'style_run': style_of_chunk,
                'parent_para': para,
                'context': paragraph_context,
                'chunk_index': chunk_index_in_para
            })
            
    total_tasks = len(tasks)
    if total_tasks == 0:
        print("未发现需要翻译的文本内容。")
        input_file_obj.seek(0)
        return input_file_obj

    print(f"共找到 {total_tasks} 个独立的文本块需要翻译。开始并发翻译...")

    # --- 阶段 2: 并发执行翻译任务 ---
    reconstruction_data = defaultdict(list)
    processed_count = 0
    MAX_WORKERS = 10

    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
        # 【请在此处选择要使用的翻译函数】
        # 使用 Azure:
        # target_translation_function = get_llm_translation
        # 使用 Xinference:
        target_translation_function = get_llm_translation_xinference

        future_to_task = {
            executor.submit(target_translation_function, task['text'], task['context'], target_language): task
            for task in tasks
        }

        for future in as_completed(future_to_task):
            task = future_to_task[future]
            parent_para = task['parent_para']
            
            try:
                translated_text = future.result()
                reconstruction_data[parent_para].append({
                    'text': translated_text,
                    'style_run': task['style_run'],
                    'chunk_index': task['chunk_index'] 
                })
            except Exception as exc:
                print(f"!!! 文本块 '{task['text'][:30]}...' 翻译生成了一个异常: {exc}")
                reconstruction_data[parent_para].append({
                    'text': task['text'], # 翻译失败，保留原文
                    'style_run': task['style_run'],
                    'chunk_index': task['chunk_index']
                })
            
            processed_count += 1
            if progress_callback:
                progress_callback(processed_count / total_tasks)

    # --- 阶段 3: 在主线程中安全地重建文档 ---
    print("所有翻译任务已完成。正在重建演示文稿...")
    
    for para, translated_chunks in reconstruction_data.items():
        # 按原始顺序排序
        sorted_chunks = sorted(translated_chunks, key=operator.itemgetter('chunk_index'))

        # 清空段落现有内容
        # 保留第一个 run 作为添加新内容的起点，其余的都移除
        if para.runs:
            first_run = para.runs[0]
            first_run.text = ''
            p = para._p
            for i in range(len(para.runs) - 1, 0, -1):
                p.remove(para.runs[i]._r)
        else:
            # 如果段落完全是空的，就添加一个 run
            first_run = para.add_run()


        # 根据排序后的chunks重建段落
        is_first_chunk = True
        for chunk in sorted_chunks:
            # 对于第一个chunk，我们重用已存在的run
            if is_first_chunk:
                new_run = first_run
                is_first_chunk = False
            else:
                # 后续的chunk创建新的run
                new_run = para.add_run()
            
            new_run.text = chunk['text']
            apply_style_to_pptx_run(chunk['style_run'], new_run)

    print("演示文稿重建完成。")
    
    output_buffer = io.BytesIO()
    prs.save(output_buffer)
    output_buffer.seek(0)
    
    return output_buffer