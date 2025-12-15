import streamlit as st
from docx import Document
from docx.shared import Pt
import io
import re

def process_docx(uploaded_file, font_size=16, add_numbering=False, fix_level_skips=False):
    """
    打开一个从 Streamlit 上传的 Word 文档文件对象，
    根据选项对标题进行处理，并返回一个包含修改后文档的内存中二进制流。

    :param uploaded_file: Streamlit file_uploader 上传的文件对象
    :param font_size: 需要设置的标题字号大小
    :param add_numbering: 是否为缺少序号的标题自动添加序号
    :param fix_level_skips: 是否自动修正标题跳级问题 (例如 H1 -> H3)
    :return: io.BytesIO 对象，包含处理后的 .docx 文件内容
    """
    try:
        document = Document(uploaded_file)
        
        counters = [0] * 9
        numbering_pattern = re.compile(r'''^\s*\d+(\.\d+)*[.\s\u3000]''')
        last_level = 0  # 跟踪上一个标题的级别

        for paragraph in document.paragraphs:
            if paragraph.style and paragraph.style.name.startswith('Heading'):
                try:
                    level = int(paragraph.style.name.split(' ')[-1])

                    # --- 自动修正标题跳级 ---
                    if fix_level_skips and level > last_level + 1:
                        new_level = last_level + 1
                        try:
                            # 尝试应用修正后的样式 (例如 'Heading 2')
                            paragraph.style = f'Heading {new_level}'
                            level = new_level # 如果成功，则更新 level
                        except KeyError:
                            # 如果文档中不存在该样式，则发出警告但继续处理
                            st.warning(f"文档缺少“Heading {new_level}”样式，无法修正标题跳级。将按原样式处理。")
                    
                    last_level = level
                    # --- 修正结束 ---

                    # --- 自动添加序号的逻辑 ---
                    if add_numbering and not numbering_pattern.match(paragraph.text):
                        if 1 <= level <= 9:
                            counters[level - 1] += 1
                            for i in range(level, 9):
                                counters[i] = 0
                            
                            numbering_parts = [str(c) for c in counters[:level] if c > 0]
                            numbering_str = ".".join(numbering_parts)
                            
                            if paragraph.runs:
                                paragraph.runs[0].text = f"{numbering_str} {paragraph.runs[0].text}"
                            else:
                                paragraph.add_run(f"{numbering_str} ")

                except (ValueError, IndexError):
                    # 如果样式名称不规范，则跳过此段落的逻辑处理
                    pass
                
                # --- 统一修改标题格式的逻辑 ---
                for run in paragraph.runs:
                    run.font.size = Pt(font_size)
                    run.font.italic = False
        
        file_stream = io.BytesIO()
        document.save(file_stream)
        file_stream.seek(0)
        
        return file_stream

    except Exception as e:
        st.error(f"处理文档时发生错误: {e}")
        st.error("请确保您上传的是一个有效的 .docx 文件。")
        return None

# --- Streamlit 界面布局 ---

st.title("✍️ Word 文档标题格式批量修改器")

st.markdown("""
这是一个多功能小工具，可以帮你自动化处理 Word 文档中的标题格式：
- **统一字号**: 将所有级别的标题设置为你想要的字号。
- **取消斜体**: 自动将所有标题的斜体样式去除。
- **自动编号**: 为没有序号的标题添加层级编号 (例如 1. / 1.1 / 1.1.1)。
- **修正跳级**: 自动修复不连续的标题级别 (例如一级标题后直接是三级，会自动修正为二级)。
""")

uploaded_file = st.file_uploader(
    "请在此处上传您的 .docx 文档",
    type=['docx'],
    help="请上传 .docx 格式的 Word 文档"
)

if uploaded_file is not None:
    st.success(f"文件上传成功: **{uploaded_file.name}**")
    st.divider()

    # --- 功能选项 ---
    font_size_to_set = st.number_input(
        "设置标题字号:", 
        min_value=1, 
        max_value=100, 
        value=16, 
        step=1,
        help="所有检测到的标题都会被设置为这个字号"
    )

    col1, col2 = st.columns(2)
    with col1:
        add_numbering = st.toggle(
            "为标题自动编号",
            value=True,
            help="为“标题”样式且开头没有序号的段落，添加多级序号。"
        )
    with col2:
        fix_level_skips = st.toggle(
            "自动修正标题跳级",
            value=True,
            help="例如，当一级标题后直接是三级标题时，会自动将其修正为二级标题，以保证层级连续。"
        )

    # --- 处理与下载 ---
    if st.button("🚀 开始处理并生成下载文件", type="primary"):
        
        tasks = []
        if add_numbering:
            tasks.append("自动编号")
        if fix_level_skips:
            tasks.append("修正跳级")
        
        spinner_text = f"处理中... 正在修改字号、取消斜体"
        if tasks:
            spinner_text += f"，并执行: {', '.join(tasks)}"

        with st.spinner(spinner_text + "..."):
            processed_stream = process_docx(
                uploaded_file, 
                font_size_to_set, 
                add_numbering=add_numbering,
                fix_level_skips=fix_level_skips
            )
            
            if processed_stream:
                st.success("🎉 文档处理完成！")
                
                new_file_name = f"修改后_{uploaded_file.name}"
                
                st.download_button(
                    label="📥 点击下载修改后的文档",
                    data=processed_stream,
                    file_name=new_file_name,
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )
