import io
import operator
import logging
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from typing import Callable, Optional, List, Dict, Any, Union

# Third-party libraries
try:
    from docx import Document
    from docx.text.run import Run
except ImportError:
    Document = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    from pptx import Presentation
    from pptx.text.text import _Run
except ImportError:
    Presentation = None

logger = logging.getLogger(__name__)

class DocxParser:
    """
    Parser for .docx files.
    Extracts text chunks while preserving styles, delegates translation via callback, and rebuilds the document.
    """
    
    @staticmethod
    def get_run_style(run: 'Run'):
        """Get style attributes of a run."""
        font = run.font
        return (font.name, font.size, font.bold, font.italic, font.underline, 
                font.color.rgb if font.color and font.color.rgb else None, 
                font.highlight_color)

    @staticmethod
    def apply_style_to_run(source_run: 'Run', target_run: 'Run'):
        """Apply style from source run to target run."""
        target_run.font.name = source_run.font.name
        target_run.font.size = source_run.font.size
        target_run.font.bold = source_run.font.bold
        target_run.font.italic = source_run.font.italic
        target_run.font.underline = source_run.font.underline
        if source_run.font.color and source_run.font.color.rgb:
            target_run.font.color.rgb = source_run.font.color.rgb
        if source_run.font.highlight_color:
            target_run.font.highlight_color = source_run.font.highlight_color

    def process(self, input_file_obj: io.BytesIO, translate_callback: Callable[[str, str], str], max_workers: int = 10) -> io.BytesIO:
        """
        Process the docx file: parse -> translate (concurrently) -> rebuild.
        
        Args:
            input_file_obj: Input file stream.
            translate_callback: Function(text, context) -> translated_text.
            max_workers: Number of threads for concurrent translation.
            
        Returns:
            io.BytesIO: The translated file stream.
        """
        if not Document:
            raise ImportError("python-docx is not installed.")

        doc = Document(input_file_obj)
        
        # --- Phase 1: Collect tasks ---
        tasks = []
        all_paragraphs = list(doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    all_paragraphs.extend(cell.paragraphs)

        logger.info(f"Found {len(all_paragraphs)} paragraphs. Analyzing and chunking...")

        for para in all_paragraphs:
            paragraph_context = para.text
            if not para.runs or not paragraph_context.strip():
                continue
            
            chunk_index_in_para = 0
            current_text_chunk = ""
            style_of_chunk = para.runs[0]

            for run in para.runs:
                if self.get_run_style(run) == self.get_run_style(style_of_chunk):
                    current_text_chunk += run.text
                else:
                    if current_text_chunk.strip():
                        tasks.append({
                            'text': current_text_chunk,
                            'style_run': style_of_chunk,
                            'parent_para': para,
                            'context': paragraph_context,
                            'chunk_index': chunk_index_in_para
                        })
                        chunk_index_in_para += 1
                    current_text_chunk = run.text
                    style_of_chunk = run
            
            if current_text_chunk.strip():
                tasks.append({
                    'text': current_text_chunk,
                    'style_run': style_of_chunk,
                    'parent_para': para,
                    'context': paragraph_context,
                    'chunk_index': chunk_index_in_para
                })

        total_tasks = len(tasks)
        if total_tasks == 0:
            logger.info("No translatable content found.")
            input_file_obj.seek(0)
            return input_file_obj

        logger.info(f"Found {total_tasks} text chunks. Starting concurrent translation...")

        # --- Phase 2: Execute translation ---
        reconstruction_data = defaultdict(list)
        processed_count = 0

        with ThreadPoolExecutor(max_workers=max_workers) as executor:
            future_to_task = {
                executor.submit(translate_callback, task['text'], task['context']): task
                for task in tasks
            }

            for future in as_completed(future_to_task):
                task = future_to_task[future]
                parent_para = task['parent_para']
                
                try:
                    translated_text = future.result()
                    reconstruction_data[parent_para].append({
                        'text': translated_text,
                        'style_run': task['style_run'],
                        'chunk_index': task['chunk_index'] 
                    })
                except Exception as exc:
                    logger.error(f"Error translating chunk '{task['text'][:30]}...': {exc}")
                    reconstruction_data[parent_para].append({
                        'text': task['text'], # Keep original on failure
                        'style_run': task['style_run'],
                        'chunk_index': task['chunk_index']
                    })
                
                processed_count += 1
                if processed_count % 10 == 0 or processed_count == total_tasks:
                    logger.info(f"Progress: {processed_count}/{total_tasks}")

        # --- Phase 3: Rebuild document ---
        logger.info("Rebuilding document...")
        
        for para, translated_chunks in reconstruction_data.items():
            sorted_chunks = sorted(translated_chunks, key=operator.itemgetter('chunk_index'))

            # Clear existing content
            p = para._p
            p.clear_content()

            # Rebuild
            for chunk in sorted_chunks:
                new_run = para.add_run(chunk['text'])
                self.apply_style_to_run(chunk['style_run'], new_run)

        logger.info("Document reconstruction complete.")
        
        output_buffer = io.BytesIO()
        doc.save(output_buffer)
        output_buffer.seek(0)
        
        return output_buffer


class XlsxParser:
    """
    Parser for .xlsx files.
    Handles cell translation and optionally sheet name translation.
    """

    def process(self, input_file_obj: io.BytesIO, translate_callback: Callable[[str], str], 
                translate_sheet_names: bool = False, max_workers: int = 10) -> io.BytesIO:
        """
        Process the xlsx file.
        
        Args:
            input_file_obj: Input file stream.
            translate_callback: Function(text) -> translated_text (Note: context is not typically used for cells here, but can be added if needed).
            translate_sheet_names: Whether to translate sheet names.
            max_workers: Number of threads.
        """
        if not openpyxl:
            raise ImportError("openpyxl is not installed.")

        workbook = openpyxl.load_workbook(input_file_obj)
        tasks = []
        
        # 1. Collect tasks
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.data_type == 's' and cell.value and isinstance(cell.value, str) and cell.value.strip():
                        tasks.append({
                            'type': 'cell',
                            'text': cell.value,
                            'target_object': cell
                        })
                        
        if translate_sheet_names:
            logger.warning("Sheet name translation enabled. Formula links might be broken.")
            for sheet in workbook.worksheets:
                if sheet.title and sheet.title.strip():
                    tasks.append({
                        'type': 'sheet',
                        'text': sheet.title,
                        'target_object': sheet
                    })

        total_tasks = len(tasks)
        if total_tasks == 0:
            logger.info("No translatable content found.")
            input_file_obj.seek(0)
            return input_file_obj

        logger.info(f"Found {total_tasks} items to translate.")

        # 2. Execute translation
        processed_count = 0
        with ThreadPoolExecutor(max_workers=max_workers) as executor:
            future_to_task = {
                executor.submit(translate_callback, task['text']): task
                for task in tasks
            }

            for future in as_completed(future_to_task):
                task = future_to_task[future]
                try:
                    translated_text = future.result()
                    
                    if task['type'] == 'cell':
                        task['target_object'].value = translated_text
                    elif task['type'] == 'sheet':
                        task['target_object'].title = translated_text
                        
                except Exception as exc:
                    logger.error(f"Error translating item '{task['text'][:30]}...' ({task['type']}): {exc}")
                
                processed_count += 1
                if processed_count % 50 == 0 or processed_count == total_tasks:
                     logger.info(f"Progress: {processed_count}/{total_tasks}")

        logger.info("XLSX translation complete.")
        
        output_buffer = io.BytesIO()
        workbook.save(output_buffer)
        output_buffer.seek(0)
        
        return output_buffer


class PptxParser:
    """
    Parser for .pptx files.
    Extracts text from slides/shapes/tables, preserving styles.
    """
    
    @staticmethod
    def get_run_style(run: '_Run'):
        """Get style attributes of a PPTX run."""
        font = run.font
        color_info = None
        if hasattr(font.color, 'type') and font.color.type is not None:
            if font.color.type == 1: # MSO_COLOR_TYPE.RGB
                color_info = ('RGB', font.color.rgb)
            elif font.color.type == 2: # MSO_COLOR_TYPE.SCHEME
                color_info = ('SCHEME', font.color.theme_color, font.color.brightness)

        return (
            font.name, font.size, font.bold, font.italic, font.underline, color_info
        )

    @staticmethod
    def apply_style_to_run(style_run: '_Run', target_run: '_Run'):
        """Apply style from source run to target run."""
        s_font = style_run.font
        t_font = target_run.font
        
        t_font.name = s_font.name
        t_font.size = s_font.size
        t_font.bold = s_font.bold
        t_font.italic = s_font.italic
        t_font.underline = s_font.underline

        s_color = s_font.color
        t_color = t_font.color
        if hasattr(s_color, 'type') and s_color.type is not None:
            if s_color.type == 1: # MSO_COLOR_TYPE.RGB
                t_color.rgb = s_color.rgb
            elif s_color.type == 2: # MSO_COLOR_TYPE.SCHEME
                t_color.theme_color = s_color.theme_color
                t_color.brightness = s_color.brightness

    def process(self, input_file_obj: io.BytesIO, translate_callback: Callable[[str, str], str], max_workers: int = 10) -> io.BytesIO:
        """
        Process the pptx file.
        
        Args:
            input_file_obj: Input file stream.
            translate_callback: Function(text, context) -> translated_text.
            max_workers: Number of threads.
        """
        if not Presentation:
            raise ImportError("python-pptx is not installed.")

        prs = Presentation(input_file_obj)
        
        # --- Phase 1: Collect tasks ---
        tasks = []
        all_paragraphs = []

        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    all_paragraphs.extend(shape.text_frame.paragraphs)
                if shape.has_table:
                    for row in shape.table.rows:
                        for cell in row.cells:
                            all_paragraphs.extend(cell.text_frame.paragraphs)

        logger.info(f"Found {len(all_paragraphs)} paragraphs. Analyzing and chunking...")

        for para in all_paragraphs:
            paragraph_context = para.text
            if not para.runs or not paragraph_context.strip():
                continue
            
            chunk_index_in_para = 0
            current_text_chunk = ""
            style_of_chunk = para.runs[0]

            for run in para.runs:
                if self.get_run_style(run) == self.get_run_style(style_of_chunk):
                    current_text_chunk += run.text
                else:
                    if current_text_chunk.strip():
                        tasks.append({
                            'text': current_text_chunk,
                            'style_run': style_of_chunk,
                            'parent_para': para,
                            'context': paragraph_context,
                            'chunk_index': chunk_index_in_para
                        })
                        chunk_index_in_para += 1
                    current_text_chunk = run.text
                    style_of_chunk = run
            
            if current_text_chunk.strip():
                tasks.append({
                    'text': current_text_chunk,
                    'style_run': style_of_chunk,
                    'parent_para': para,
                    'context': paragraph_context,
                    'chunk_index': chunk_index_in_para
                })
                
        total_tasks = len(tasks)
        if total_tasks == 0:
            logger.info("No translatable content found.")
            input_file_obj.seek(0)
            return input_file_obj

        logger.info(f"Found {total_tasks} text chunks. Starting concurrent translation...")

        # --- Phase 2: Execute translation ---
        reconstruction_data = defaultdict(list)
        processed_count = 0

        with ThreadPoolExecutor(max_workers=max_workers) as executor:
            future_to_task = {
                executor.submit(translate_callback, task['text'], task['context']): task
                for task in tasks
            }

            for future in as_completed(future_to_task):
                task = future_to_task[future]
                parent_para = task['parent_para']
                
                try:
                    translated_text = future.result()
                    reconstruction_data[parent_para].append({
                        'text': translated_text,
                        'style_run': task['style_run'],
                        'chunk_index': task['chunk_index'] 
                    })
                except Exception as exc:
                    logger.error(f"Error translating chunk '{task['text'][:30]}...': {exc}")
                    reconstruction_data[parent_para].append({
                        'text': task['text'],
                        'style_run': task['style_run'],
                        'chunk_index': task['chunk_index']
                    })
                
                processed_count += 1
                if processed_count % 10 == 0 or processed_count == total_tasks:
                    logger.info(f"Progress: {processed_count}/{total_tasks}")

        # --- Phase 3: Rebuild ---
        logger.info("Rebuilding presentation...")
        
        for para, translated_chunks in reconstruction_data.items():
            sorted_chunks = sorted(translated_chunks, key=operator.itemgetter('chunk_index'))

            # Clear content: preserve first run for property inheritance if possible, remove others
            if para.runs:
                first_run = para.runs[0]
                first_run.text = ''
                p = para._p
                # Remove runs from end to start (except first)
                for i in range(len(para.runs) - 1, 0, -1):
                    p.remove(para.runs[i]._r)
            else:
                first_run = para.add_run()

            # Rebuild
            is_first_chunk = True
            for chunk in sorted_chunks:
                if is_first_chunk:
                    new_run = first_run
                    is_first_chunk = False
                else:
                    new_run = para.add_run()
                
                new_run.text = chunk['text']
                self.apply_style_to_run(chunk['style_run'], new_run)

        logger.info("Presentation reconstruction complete.")
        
        output_buffer = io.BytesIO()
        prs.save(output_buffer)
        output_buffer.seek(0)
        
        return output_buffer
